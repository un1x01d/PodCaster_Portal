import pandas as pd
import numpy as np
import os
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

output_file = "heavy_enterprise_data.xlsx"
num_rows = 200000 # Adjusted for size target

print(f"Generating {num_rows} rows of data...")

# Tab 1: Sales Data
df_sales = pd.DataFrame({
    "Transaction_ID": [f"TXN-{i:07d}" for i in range(num_rows)],
    "Date": pd.date_range(start="2023-01-01", periods=num_rows, freq="min"),
    "Store_ID": np.random.randint(100, 200, size=num_rows),
    "Product_SKU": [f"SKU-{np.random.randint(1000, 9999)}" for _ in range(num_rows)],
    "Quantity": np.random.randint(1, 10, size=num_rows),
    "Unit_Price": np.random.uniform(10.0, 500.0, size=num_rows).round(2),
    "Tax_Rate": 0.08,
    "Discount": np.random.uniform(0, 50.0, size=num_rows).round(2),
    "Region": np.random.choice(["North", "South", "East", "West", "Central"], size=num_rows),
    "Customer_Segment": np.random.choice(["Enterprise", "SMB", "Consumer", "Government"], size=num_rows),
    "Payment_Method": np.random.choice(["Credit Card", "Wire Transfer", "Net 30", "Cryptocurrency"], size=num_rows),
    "Is_Flagged": np.random.choice([True, False], size=num_rows),
    "Processing_Time_Sec": np.random.uniform(0.1, 5.0, size=num_rows).round(3)
})

# Add dummy columns to bloat size
for i in range(1, 21):
    df_sales[f"Meta_Field_{i}"] = [f"Value_String_Long_Data_Padding_{np.random.randint(1000, 9999)}" for _ in range(num_rows)]

# Tab 2: Reference Data
df_refs = pd.DataFrame({
    "Store_ID": np.arange(100, 200),
    "Store_Location": [f"City_{i}" for i in range(100)],
    "Manager": [f"Manager_{i}" for i in range(100)],
    "Capacity": np.random.randint(1000, 5000, size=100)
})

print("Writing to Excel (this may take a minute)...")
with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
    df_sales.to_excel(writer, sheet_name='Sales_Data', index=False)
    df_refs.to_excel(writer, sheet_name='Reference_Master', index=False)

print("Applying formatting and formulas...")
wb = load_workbook(output_file)
ws = wb['Sales_Data']

# 1. Frozen Panes
ws.freeze_panes = "A2"

# 2. Formulas
# We'll add some formulas in new columns at the end
# Column count is Transaction_ID (1) + Date (1) + ... + Meta_Fields (20) = 13 + 20 = 33
# Let's add Gross_Total in column 34 and Net_Total in column 35
col_qty = "E"
col_price = "F"
col_tax = "G"
col_disc = "H"

ws.cell(row=1, column=34).value = "Gross_Total_Formula"
ws.cell(row=1, column=35).value = "Net_Total_Formula"

# For heavy processing, we'll only add formulas for the first 10k rows to avoid massive save times, 
# or just do a few to show it works if the goal is "challenging". 
# Actually, let's do all of them but in a smart way. 
# openpyxl is slow for row-by-row formulas. 
# Let's just do the first 5000 rows with formulas.
for r in range(2, 5002):
    ws.cell(row=r, column=34).value = f"={col_qty}{r}*{col_price}{r}"
    ws.cell(row=r, column=35).value = f"=({col_qty}{r}*{col_price}{r})*(1+{col_tax}{r})-{col_disc}{r}"

# Add a summary tab with pivot-like formulas
if 'Summary' not in wb.sheetnames:
    wb.create_sheet('Summary')
ws_sum = wb['Summary']
ws_sum['A1'] = "Total Quantity Sum"
ws_sum['B1'] = f"=SUM(Sales_Data!{col_qty}:{col_qty})"
ws_sum['A2'] = "Average Unit Price"
ws_sum['B2'] = f"=AVERAGE(Sales_Data!{col_price}:{col_price})"

print("Saving final file...")
wb.save(output_file)

size_mb = os.path.getsize(output_file) / (1024 * 1024)
print(f"Done! Created '{output_file}' ({size_mb:.2f} MB)")

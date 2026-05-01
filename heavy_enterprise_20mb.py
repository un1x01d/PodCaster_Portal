import pandas as pd
import numpy as np
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, timedelta

output_file = "Corporate Financial Ledger 2020 2026.xlsx"
# Targeted row count for ~20MB based on previous metrics
num_rows = 160000 

print(f"Re-evaluating the budget... generating {num_rows} records for the 20MB reporting target.")

# --- ACCOUNTING DATA POOLS ---
regions = ["North America", "European Union", "Asia Pacific", "Latin America", "Middle East"]
countries = {
    "North America": ["United States", "Canada", "Mexico"],
    "European Union": ["United Kingdom", "Germany", "France", "Spain"],
    "Asia Pacific": ["Singapore", "Japan", "Australia", "India"],
    "Latin America": ["Brazil", "Argentina", "Chile"],
    "Middle East": ["UAE", "Saudi Arabia", "Qatar"]
}
segments = ["Enterprise Tier", "Mid Market", "Small Business", "Public Sector"]
channels = ["Direct Field Sales", "Global Partner Network", "Digital Storefront", "Authorized Resellers"]
product_catalog = [
    {"name": "Quantum Mainframe X1", "cat": "Hard Assets", "vendor": "Stellar Corp", "price": 45000.0, "cost": 28500.0},
    {"name": "Nebula Storage Array", "cat": "Hard Assets", "vendor": "Stellar Corp", "price": 12000.0, "cost": 7200.0},
    {"name": "Cyber Shield Premium", "cat": "Software Licenses", "vendor": "Nexus Security", "price": 3500.0, "cost": 450.0},
    {"name": "Data Flow Analytics", "cat": "Software Licenses", "vendor": "Nexus Security", "price": 1800.0, "cost": 220.0},
    {"name": "Enterprise Cloud Migration", "cat": "Professional Services", "vendor": "Prime Consulting", "price": 75000.0, "cost": 48000.0},
    {"name": "Quarterly Security Audit", "cat": "Professional Services", "vendor": "Prime Consulting", "price": 15000.0, "cost": 9500.0},
]

wb = Workbook()

# 1. Transaction Ledger
ws_main = wb.active
ws_main.title = "Transaction Ledger"

# 2. Product Inventory
ws_prod = wb.create_sheet("Product Inventory")
# 3. Entity Hierarchy
ws_org = wb.create_sheet("Entity Hierarchy")
# 4. Operational Expenses
ws_opex = wb.create_sheet("Operational Expenses")
# 5. Executive Summary
ws_sum = wb.create_sheet("Executive Summary")

# --- STYLING ---
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

# --- LOOKUPS ---
ws_prod.append(["Product Name", "Asset Class", "Prime Vendor", "Standard List Price", "Internal Unit Cost", "Warranty Period", "Logistics Lead Time", "Tax Status"])
for p in product_catalog:
    ws_prod.append([p["name"], p["cat"], p["vendor"], p["price"], p["cost"], "36 Months", "14 Days", "Taxable"])

ws_org.append(["Business Region", "Country Office", "Senior Controller", "Employee Headcount", "Reporting Currency"])
for reg, count_list in countries.items():
    for country in count_list:
        ws_org.append([reg, country, f"Controller {country}", np.random.randint(100, 1500), "USD"])

# --- MAIN LEDGER ---
headers = [
    "Posting Date", "Journal Entry ID", "Business Region", "Entity Country", "Customer Segment", 
    "Sales Channel", "Asset Class", "Product Model", "Vendor Name", "Unit Quantity", 
    "Unit List Price", "Unit Production Cost", "Applied Discount Rate", "Gross Revenue", "Discount Value", 
    "Net Sales Revenue", "Cost of Goods Sold", "Operating Gross Profit", "Performance Margin", "Internal Audit Notes"
]
ws_main.append(headers)

for cell in ws_main[1]:
    cell.font = header_font
    cell.fill = header_fill

start_date = datetime(2020, 1, 1)
end_date = datetime(2026, 12, 31)
delta_days = (end_date - start_date).days

print(f"Streaming {num_rows} transactions...")
for i in range(num_rows):
    progress = i / num_rows
    curr_date = start_date + timedelta(days=int(progress * delta_days))
    
    region = np.random.choice(regions)
    country = np.random.choice(countries[region])
    product = np.random.choice(product_catalog)
    qty = np.random.randint(1, 200)
    price = product["price"] * np.random.uniform(0.98, 1.02)
    cost = product["cost"] * np.random.uniform(0.99, 1.01)
    discount = round(np.random.beta(1, 10) * 0.12, 4)
    
    r_idx = i + 2
    row = [
        curr_date, f"GL-{(i+1000000)}", region, country, np.random.choice(segments),
        np.random.choice(channels), product["cat"], product["name"], product["vendor"], qty,
        round(price, 2), round(cost, 2), discount
    ]
    ws_main.append(row)
    
    # Formulas
    ws_main.cell(row=r_idx, column=14).value = f"=J{r_idx}*K{r_idx}"
    ws_main.cell(row=r_idx, column=15).value = f"=N{r_idx}*M{r_idx}"
    ws_main.cell(row=r_idx, column=16).value = f"=N{r_idx}-O{r_idx}"
    ws_main.cell(row=r_idx, column=17).value = f"=J{r_idx}*L{r_idx}"
    ws_main.cell(row=r_idx, column=18).value = f"=P{r_idx}-Q{r_idx}"
    ws_main.cell(row=r_idx, column=19).value = f"=IF(P{r_idx}=0,0,R{r_idx}/P{r_idx})"
    ws_main.cell(row=r_idx, column=20).value = f"Audit verified for batch {i//5000}. No discrepancies noted."

ws_main.freeze_panes = "A2"

# --- OpEx ---
for year in range(2020, 2027):
    for month in range(1, 13):
        for reg in regions:
            ws_opex.append([year, month, reg, "Operations", 100000, 95000, "=F{0}-E{0}".format(ws_opex.max_row+1)])

# --- Summary ---
ws_sum.append(["Fiscal Year", "Net Revenue", "Gross Profit", "Avg Margin"])
for year in range(2020, 2027):
    r = ws_sum.max_row + 1
    ws_sum.cell(row=r, column=1).value = year
    ws_sum.cell(row=r, column=2).value = f"=SUMIFS('Transaction Ledger'!P:P, 'Transaction Ledger'!A:A, \">=\"&DATE({year},1,1), 'Transaction Ledger'!A:A, \"<=\"&DATE({year},12,31))"
    ws_sum.cell(row=r, column=3).value = f"=SUMIFS('Transaction Ledger'!R:R, 'Transaction Ledger'!A:A, \">=\"&DATE({year},1,1), 'Transaction Ledger'!A:A, \"<=\"&DATE({year},12,31))"
    ws_sum.cell(row=r, column=4).value = f"=AVERAGEIFS('Transaction Ledger'!S:S, 'Transaction Ledger'!A:A, \">=\"&DATE({year},1,1), 'Transaction Ledger'!A:A, \"<=\"&DATE({year},12,31))"

wb.save(output_file)
print(f"Final size: {os.path.getsize(output_file)/(1024*1024):.2f} MB")

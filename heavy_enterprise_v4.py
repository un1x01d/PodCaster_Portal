import pandas as pd
import numpy as np
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

output_file = "enterprise_ops_2026_final.xlsx"
# 250k rows with detailed notes should hit the 20MB sweet spot
num_rows = 250000 

print(f"Generating {num_rows} rows of high-fidelity enterprise data...")

# Business Context
regions = ["North America", "EMEA", "APAC", "LATAM", "DACH", "Nordics"]
sectors = ["Healthcare", "FinTech", "Energy", "Logistics", "Aerospace", "Telco"]
vendors = ["Global Dynamics", "Stellar Systems", "Prime Logistics", "Nexus Energy", "Quantum Tech"]

wb = Workbook()
ws = wb.active
ws.title = "Dataset"

headers = ["Order_ID", "Fiscal_Period", "Region", "Sector", "Vendor", "Qty", "Unit_Price", "Disc_Pct", "Tax_Pct", "Status", "Audit_Notes", "Gross_Total", "Net_Total"]
ws.append(headers)

# Use bold headers
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")

print("Streaming data to workbook...")
for i in range(num_rows):
    region = regions[i % len(regions)]
    sector = sectors[i % len(sectors)]
    vendor = vendors[i % len(vendors)]
    
    qty = (i % 50) + 1
    unit_price = 1000.0 + (i % 5000)
    discount = 0.05 + (0.01 * (i % 15))
    tax = 0.12
    
    # Realistic but long Audit Notes to increase file size naturally
    notes = f"Transaction {i} verified for {sector} in {region}. Compliance check by {vendor} internal audit division. Document ID: {i*7}."
    
    r_idx = i + 2
    row_data = [
        f"ORD-2026-{i:07d}", "2026-Q1", region, sector, vendor,
        qty, unit_price, discount, tax, "Verified", notes
    ]
    ws.append(row_data)
    
    # Add Formulas
    # Gross_Total (L) = F * G
    ws.cell(row=r_idx, column=12).value = f"=F{r_idx}*G{r_idx}"
    # Net_Total (M) = (F*G)*(1-H)*(1+I)
    ws.cell(row=r_idx, column=13).value = f"=(F{r_idx}*G{r_idx})*(1-H{r_idx})*(1+I{r_idx})"

ws.freeze_panes = "A2"

print("Saving (High-volume formula save)...")
wb.save(output_file)

size_mb = os.path.getsize(output_file) / (1024 * 1024)
print(f"Final File: '{output_file}' ({size_mb:.2f} MB)")

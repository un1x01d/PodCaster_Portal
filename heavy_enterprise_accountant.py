import pandas as pd
import numpy as np
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime, timedelta

output_file = "Global Enterprise Ledger 2020 2026.xlsx"
# Increased row count to ensure we hit the ~30MB mark after XLSX compression
num_rows = 300000 

print(f"Opening the books... generating {num_rows} line items for the audit.")

# --- ACCOUNTING DATA POOLS ---
regions = ["North America", "European Union", "Asia Pacific", "Latin America", "Middle East"]
countries = {
    "North America": ["United States", "Canada", "Mexico"],
    "European Union": ["United Kingdom", "Germany", "France", "Italy"],
    "Asia Pacific": ["Singapore", "Japan", "Australia", "India"],
    "Latin America": ["Brazil", "Argentina", "Chile"],
    "Middle East": ["UAE", "Saudi Arabia", "Qatar"]
}
segments = ["Enterprise Tier", "Mid Market", "Small Business", "Public Sector"]
channels = ["Direct Field Sales", "Global Partner Network", "Digital Storefront", "Authorized Resellers"]
product_catalog = [
    {"name": "Quantum Mainframe X1", "cat": "Hard Assets", "vendor": "Stellar Corp", "price": 45000.0, "cost": 28500.0},
    {"name": "Nebula Storage Array", "cat": "Hard Assets", "vendor": "Stellar Corp", "price": 12000.0, "cost": 7200.0},
    {"name": "Cyber Shield Premium", "cat": "Software Licenses", "vendor": "Nexus Security", "price": 3500.0, "cost": 450.0},
    {"name": "Data Flow Analytics", "cat": "Software Licenses", "vendor": "Nexus Security", "price": 1800.0, "cost": 220.0},
    {"name": "Enterprise Cloud Migration", "cat": "Professional Services", "vendor": "Prime Consulting", "price": 75000.0, "cost": 48000.0},
    {"name": "Quarterly Security Audit", "cat": "Professional Services", "vendor": "Prime Consulting", "price": 15000.0, "cost": 9500.0},
]

wb = Workbook()

# 1. Transaction Ledger (Main Fact Table)
ws_main = wb.active
ws_main.title = "Transaction Ledger"

# 2. Product Inventory
ws_prod = wb.create_sheet("Product Inventory")
# 3. Entity Hierarchy
ws_org = wb.create_sheet("Entity Hierarchy")
# 4. Operational Expenses
ws_opex = wb.create_sheet("Operational Expenses")
# 5. Executive Summary
ws_sum = wb.create_sheet("Executive Summary")

# --- STYLING UTILS ---
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
center_align = Alignment(horizontal="center", vertical="center")

# --- POPULATE LOOKUPS ---
# Tab 2: Product Inventory
ws_prod.append(["Product Name", "Asset Class", "Prime Vendor", "Standard List Price", "Internal Unit Cost", "Warranty Period", "Logistics Lead Time", "Tax Status"])
for p in product_catalog:
    ws_prod.append([p["name"], p["cat"], p["vendor"], p["price"], p["cost"], "36 Months", "14 Days", "Taxable"])

# Tab 3: Entity Hierarchy
ws_org.append(["Business Region", "Country Office", "Senior Controller", "Employee Headcount", "Reporting Currency"])
for reg, count_list in countries.items():
    for country in count_list:
        ws_org.append([reg, country, f"Controller {country}", np.random.randint(100, 1500), "USD"])

# --- POPULATE TRANSACTION LEDGER ---
# 20 Professional Columns
headers = [
    "Posting Date", "Journal Entry ID", "Business Region", "Entity Country", "Customer Segment", 
    "Sales Channel", "Asset Class", "Product Model", "Vendor Name", "Unit Quantity", 
    "Unit List Price", "Unit Production Cost", "Applied Discount Rate", "Gross Revenue", "Discount Value", 
    "Net Sales Revenue", "Cost of Goods Sold", "Operating Gross Profit", "Performance Margin", "Internal Audit Notes"
]
ws_main.append(headers)

# Apply Styles to Main Header
for cell in ws_main[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align

start_date = datetime(2020, 1, 1)
end_date = datetime(2026, 12, 31)
delta_days = (end_date - start_date).days

print("Recording journal entries (300,000 rows)...")
for i in range(num_rows):
    # Sequential date progression for fiscal timeline
    progress = i / num_rows
    current_days = int(progress * delta_days)
    curr_date = start_date + timedelta(days=current_days)
    
    journal_id = f"GL-{(2026 - curr_date.year)}-{3000000 + i}"
    region = np.random.choice(regions)
    country = np.random.choice(countries[region])
    segment = np.random.choice(segments)
    channel = np.random.choice(channels)
    product = np.random.choice(product_catalog)
    
    qty = np.random.randint(1, 250)
    # Market variance for realism
    price = product["price"] * np.random.uniform(0.97, 1.03)
    cost = product["cost"] * np.random.uniform(0.99, 1.01)
    discount = round(np.random.beta(1, 10) * 0.15, 4)
    
    # Professional Audit Notes to ensure cell density and file size
    notes = f"Verified by internal audit. Transaction consistent with {segment} revenue recognition policy for {country} region. Batch ID {i // 1000}."
    
    r_idx = i + 2
    row = [
        curr_date, journal_id, region, country, segment,
        channel, product["cat"], product["name"], product["vendor"], qty,
        round(price, 2), round(cost, 2), discount
    ]
    ws_main.append(row)
    
    # FORMULAS (Columns 14-19)
    # N: Gross Revenue = J * K
    ws_main.cell(row=r_idx, column=14).value = f"=J{r_idx}*K{r_idx}"
    # O: Discount Value = N * M
    ws_main.cell(row=r_idx, column=15).value = f"=N{r_idx}*M{r_idx}"
    # P: Net Sales Revenue = N - O
    ws_main.cell(row=r_idx, column=16).value = f"=N{r_idx}-O{r_idx}"
    # Q: Cost of Goods Sold = J * L
    ws_main.cell(row=r_idx, column=17).value = f"=J{r_idx}*L{r_idx}"
    # R: Operating Gross Profit = P - Q
    ws_main.cell(row=r_idx, column=18).value = f"=P{r_idx}-Q{r_idx}"
    # S: Performance Margin = R / P
    ws_main.cell(row=r_idx, column=19).value = f"=IF(P{r_idx}=0,0,R{r_idx}/P{r_idx})"
    
    # T: Internal Audit Notes
    ws_main.cell(row=r_idx, column=20).value = notes
    
    if i % 50000 == 0:
        print(f"Audit progress: {i} records reconciled...")

ws_main.freeze_panes = "A2"

# --- Tab 4: Operational Expenses ---
print("Allocating operational budgets...")
ws_opex.append(["Fiscal Year", "Fiscal Month", "Business Region", "Expense Category", "Approved Budget", "Actual Expenditure", "Budget Variance"])
for year in range(2020, 2027):
    for month in range(1, 13):
        for reg in regions:
            for cat in ["Global Marketing", "Research and Development", "Corporate Facilities", "Travel and Entertainment"]:
                budget = np.random.uniform(50000, 250000)
                actual = budget * np.random.uniform(0.85, 1.15)
                row_idx_opex = ws_opex.max_row + 1
                ws_opex.append([year, month, reg, cat, round(budget, 2), round(actual, 2)])
                ws_opex.cell(row=row_idx_opex, column=7).value = f"=F{row_idx_opex}-E{row_idx_opex}"

# --- Tab 5: Executive Summary ---
print("Consolidating fiscal summary...")
ws_sum.append(["Reporting Fiscal Year", "Total Net Sales Revenue", "Total Operating Gross Profit", "Average Performance Margin"])
for year in range(2020, 2027):
    row_idx = ws_sum.max_row + 1
    ws_sum.cell(row=row_idx, column=1).value = year
    # Summing 'Transaction Ledger'!P (Net Revenue) based on Date in 'Transaction Ledger'!A
    ws_sum.cell(row=row_idx, column=2).value = f"=SUMIFS('Transaction Ledger'!P:P, 'Transaction Ledger'!A:A, \">=\"&DATE({year},1,1), 'Transaction Ledger'!A:A, \"<=\"&DATE({year},12,31))"
    ws_sum.cell(row=row_idx, column=3).value = f"=SUMIFS('Transaction Ledger'!R:R, 'Transaction Ledger'!A:A, \">=\"&DATE({year},1,1), 'Transaction Ledger'!A:A, \"<=\"&DATE({year},12,31))"
    ws_sum.cell(row=row_idx, column=4).value = f"=AVERAGEIFS('Transaction Ledger'!S:S, 'Transaction Ledger'!A:A, \">=\"&DATE({year},1,1), 'Transaction Ledger'!A:A, \"<=\"&DATE({year},12,31))"

print("Closing the books and finalizing the report...")
wb.save(output_file)

size_mb = os.path.getsize(output_file) / (1024 * 1024)
print(f"Success! Final Audit Report: '{output_file}' ({size_mb:.2f} MB)")

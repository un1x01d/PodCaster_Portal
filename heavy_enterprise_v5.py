import pandas as pd
import numpy as np
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, timedelta

output_file = "enterprise_data_2026_realistic.xlsx"
num_rows = 180000 # Targeted for ~20MB

print(f"Generating {num_rows} rows of high-fidelity data...")

# Realistic Data Pools
data_map = {
    "North America": {
        "countries": ["USA", "Canada", "Mexico"],
        "sectors": {
            "Healthcare": ["MRI Scanner", "Patient Monitor", "Ventilator"],
            "Tech": ["Cloud Server", "Firewall Appliance", "Workstation"],
            "Retail": ["Self-Checkout Kiosk", "Inventory Robot", "POS Terminal"]
        }
    },
    "EMEA": {
        "countries": ["UK", "Germany", "France", "UAE"],
        "sectors": {
            "Energy": ["Solar Array", "Wind Turbine Blade", "Smart Meter"],
            "FinTech": ["Payment Gateway", "Trading Terminal", "ATM Unit"],
            "Manufacturing": ["Robotic Arm", "CNC Machine", "Conveyor System"]
        }
    },
    "APAC": {
        "countries": ["Singapore", "Japan", "Australia", "China"],
        "sectors": {
            "Logistics": ["Warehouse Drone", "Sorting Unit", "Fleet Tracker"],
            "Consumer Electronics": ["8K Display", "Smart Hub", "Audio Processor"]
        }
    }
}

regions = list(data_map.keys())
statuses = ["Delivered", "In Transit", "Processing", "Backordered", "Cancelled"]

wb = Workbook()
ws = wb.active
ws.title = "Orders_2026"

# Define Headers
headers = [
    "Order_Date", "Order_ID", "Region", "Country", "Business_Sector", 
    "Product_Name", "Quantity", "Unit_Price", "Discount_Pct", 
    "Tax_Rate", "Gross_Revenue", "Net_Sales", "Status", "Shipping_Notes"
]
ws.append(headers)

# Formatting headers
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")

print("Writing rows...")
start_date = datetime(2026, 1, 1)

for i in range(num_rows):
    region = np.random.choice(regions)
    country = np.random.choice(data_map[region]["countries"])
    sector = np.random.choice(list(data_map[region]["sectors"].keys()))
    product = np.random.choice(data_map[region]["sectors"][sector])
    
    order_date = start_date + timedelta(days=np.random.randint(0, 365))
    order_id = f"ORD-{order_date.year}-{(100000 + i)}"
    
    qty = np.random.randint(1, 101)
    # Unit prices based on product type
    if "Scanner" in product or "Turbine" in product or "Machine" in product:
        unit_price = np.random.uniform(50000, 250000)
    elif "Drone" in product or "Server" in product:
        unit_price = np.random.uniform(5000, 15000)
    else:
        unit_price = np.random.uniform(500, 2500)
        
    discount = round(np.random.beta(2, 5) * 0.3, 3) # Realistic discount distribution
    tax = 0.05 if country == "UAE" else 0.19 if country == "Germany" else 0.08
    status = np.random.choice(statuses, p=[0.7, 0.15, 0.1, 0.03, 0.02])
    
    notes = f"Order {order_id} for {product} originated in {country} ({region}). Status: {status}."
    
    # Append row with initial values (we'll add formulas for Revenue and Sales)
    r_idx = i + 2
    row = [
        order_date, order_id, region, country, sector, 
        product, qty, round(unit_price, 2), discount, tax
    ]
    ws.append(row)
    
    # Formulas
    # K: Gross_Revenue = G * H
    ws.cell(row=r_idx, column=11).value = f"=G{r_idx}*H{r_idx}"
    # L: Net_Sales = Gross * (1-I) * (1+J)
    ws.cell(row=r_idx, column=12).value = f"=K{r_idx}*(1-I{r_idx})*(1+J{r_idx})"
    
    # Status and Notes
    ws.cell(row=r_idx, column=13).value = status
    ws.cell(row=r_idx, column=14).value = notes

# Freeze Panes
ws.freeze_panes = "A2"

print("Saving (High fidelity, 180k rows)...")
wb.save(output_file)

size_mb = os.path.getsize(output_file) / (1024 * 1024)
print(f"Final File: '{output_file}' ({size_mb:.2f} MB)")
